import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Establece el color de fondo de una celda."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=140, bottom=140, left=200, right=200):
    """Establece los márgenes internos de una celda en dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    Establece bordes en una celda.
    kwargs format: top={"sz": 12, "val": "single", "color": "336699"}
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'auto'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def create_document():
    doc = docx.Document()

    # Configuración de márgenes de página (2.5 cm = ~0.984 pulgadas)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Colores corporativos / académicos
    COLOR_PRIMARY = RGBColor(27, 54, 93)     # Azul marino institucional
    COLOR_SECONDARY = RGBColor(41, 128, 185) # Azul medio
    COLOR_DARK = RGBColor(44, 62, 80)        # Texto principal
    COLOR_GRAY = RGBColor(127, 140, 141)     # Subtítulos y metadatos

    # -------------------------------------------------------------
    # PORTADA / ENCABEZADO FORMAL
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(20)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = title_p.add_run("UNIVERSIDAD - FACULTAD DE INGENIERÍA\nINGENIERÍA DE SISTEMAS / SOFTWARE\n")
    run_inst.font.name = "Calibri"
    run_inst.font.size = Pt(13)
    run_inst.font.bold = True
    run_inst.font.color.rgb = COLOR_PRIMARY

    sub_mat = title_p.add_run("ASIGNATURA: PROCESOS DE SOFTWARE\n")
    sub_mat.font.name = "Calibri"
    sub_mat.font.size = Pt(12)
    sub_mat.font.bold = True
    sub_mat.font.color.rgb = COLOR_SECONDARY

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_line = p_line.add_run("―" * 45)
    run_line.font.color.rgb = COLOR_SECONDARY

    p_topic = doc.add_paragraph()
    p_topic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_topic.paragraph_format.space_before = Pt(12)
    p_topic.paragraph_format.space_after = Pt(6)
    run_topic = p_topic.add_run("INFORME TÉCNICO DE IMPLEMENTACIÓN")
    run_topic.font.name = "Calibri"
    run_topic.font.size = Pt(18)
    run_topic.font.bold = True
    run_topic.font.color.rgb = COLOR_PRIMARY

    p_subtopic = doc.add_paragraph()
    p_subtopic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_subtopic.paragraph_format.space_after = Pt(24)
    run_sub = p_subtopic.add_run("Automatización de Tareas en Pipeline CI con Jenkins (10% P1)\nIntegración con Repo Privado, Pruebas Automatizadas y Notificaciones")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_DARK

    # Tabla de Datos del Estudiante
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Estudiante:", "Anthonny Mosquera"),
        ("Usuario GitHub:", "AnthonnyM31"),
        ("Repositorio:", "https://github.com/AnthonnyM31/pipeline-ci-jenkins (Privado)"),
        ("Entorno de Ejecución:", "Jenkins 2.580.1 sobre Windows / Node.js v22.14.0")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        
        set_cell_background(c0, "F2F4F7")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=100, bottom=100, left=150, right=150)
        set_cell_margins(c1, top=100, bottom=100, left=150, right=150)
        
        border_spec = {"sz": 4, "val": "single", "color": "D0D5DD"}
        set_cell_border(c0, top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)
        set_cell_border(c1, top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.name = "Calibri"
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.name = "Calibri"
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_DARK

    doc.add_page_break()

    # -------------------------------------------------------------
    # FUNCIONES DE AYUDA PARA SECCIONES
    # -------------------------------------------------------------
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = COLOR_SECONDARY
        return h

    def add_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_DARK
        return p

    def add_image_placeholder(label, instructions):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F9FAFB")
        set_cell_margins(cell, top=200, bottom=200, left=250, right=250)
        border_spec = {"sz": 10, "val": "dashed", "color": "4A90E2"}
        set_cell_border(cell, top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        r_ico = p.add_run("📷 [ADJUNTAR AQUÍ CAPTURA DE PANTALLA]\n")
        r_ico.font.name = "Calibri"
        r_ico.font.bold = True
        r_ico.font.size = Pt(11)
        r_ico.font.color.rgb = RGBColor(41, 128, 185)

        r_lbl = p.add_run(f"Figura: {label}\n")
        r_lbl.font.name = "Calibri"
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(10.5)
        r_lbl.font.color.rgb = COLOR_PRIMARY

        r_inst = p.add_run(f"Instrucción: {instructions}")
        r_inst.font.name = "Calibri"
        r_inst.font.italic = True
        r_inst.font.size = Pt(9.5)
        r_inst.font.color.rgb = RGBColor(100, 110, 120)

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(8)

    def add_code_box(code_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F4F6F8")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        border_spec = {"sz": 6, "val": "single", "color": "D5DBDB"}
        set_cell_border(cell, top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(code_text)
        r.font.name = "Consolas"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(33, 47, 61)

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECCIÓN 1: INTRODUCCIÓN Y REQUISITOS
    # -------------------------------------------------------------
    add_h1("1. Introducción y Planteamiento de la Tarea")
    add_p(
        "En el marco de la asignatura Procesos de Software, la presente práctica aborda la automatización "
        "de tareas mediante Integración Continua (CI) empleando Jenkins como servidor de automatización central. "
        "El objetivo principal es implementar un flujo de trabajo estructurado en 3 etapas obligatorias, "
        "gestionando un repositorio privado en GitHub, ejecutando compilación y pruebas automatizadas, "
        "y notificando de manera bidireccional los estados de ejecución (éxito o fallo) hacia un canal externo."
    )
    add_p(
        "La orden de trabajo estipula los siguientes componentes esenciales:\n"
        "• Integración segura a un repositorio privado (Git / GitLab / GitHub).\n"
        "• Compilación y ejecución automatizada de pruebas unitarias.\n"
        "• Notificación de ejecución exitosa o fallida hacia un canal de comunicación (Discord Webhook)."
    )

    # -------------------------------------------------------------
    # SECCIÓN 2: ARQUITECTURA DEL PIPELINE (3 ETAPAS)
    # -------------------------------------------------------------
    add_h1("2. Arquitectura de la Solución y Diseño de las 3 Etapas")
    add_p(
        "Para cumplir con el estándar de la industria y la rúbrica de evaluación, se definió un Pipeline "
        "Declarativo (Jenkinsfile) modularizado exactamente en las tres etapas requeridas:"
    )

    add_h2("Etapa 1: Integración al Repositorio Privado")
    add_p(
        "En esta etapa, el agente de construcción de Jenkins se autentica contra el repositorio privado de GitHub "
        "utilizando credenciales seguras (Personal Access Token). Realiza el checkout del código en la rama 'main', "
        "e inspecciona el estado del repositorio y el hash del último commit para trazabilidad completa."
    )

    add_h2("Etapa 2: Compilación y Test Automatizados")
    add_p(
        "Esta etapa valida la integridad funcional del software. Se ejecuta una fase de build y emulación sintáctica "
        "('npm run build') seguida de la suite de 13 pruebas unitarias automatizadas con el ejecutor nativo 'node:test' "
        "('npm test'). Si alguna prueba falla, el comando devuelve un código de salida distinto de cero, deteniendo "
        "inmediatamente la tubería y marcando la etapa en estado ROJO (FAILURE)."
    )

    add_h2("Etapa 3: Notificación de Ejecución (Éxito o Fallo)")
    add_p(
        "Mediante las directivas 'post { success { ... } failure { ... } }', el pipeline garantiza la ejecución "
        "de una notificación sin importar el desenlace de las etapas previas. Utilizando un script de despacho "
        "('scripts/notify-discord.js') y la URL del Webhook cifrada en Jenkins, se envía un mensaje Embed estructurado "
        "con color verde (🟢 0x2ECC71) para aprobaciones o rojo (🔴 0xE74C3C) para rechazos, detallando el autor, "
        "el número de construcción y el enlace al log de Jenkins."
    )

    # -------------------------------------------------------------
    # SECCIÓN 3: GESTIÓN DE SEGURIDAD Y CREDENCIALES
    # -------------------------------------------------------------
    add_h1("3. Seguridad y Gestión de Credenciales en Jenkins")
    add_p(
        "Un principio fundamental en Procesos de Software y DevSecOps es la no divulgación de credenciales sensibles "
        "en texto plano. En este proyecto se garantizó que ningún token o URL de webhook se encuentre quemado en el "
        "código fuente ni en el Jenkinsfile."
    )
    add_p(
        "Se configuraron dos credenciales globales en el Administrador de Credenciales de Jenkins (Jenkins Credentials Store):\n"
        "1. 'github-credentials': Credencial de tipo 'Username with password' que almacena el usuario 'AnthonnyM31' "
        "y el Personal Access Token (PAT) clásico con permisos 'repo' otorgados por GitHub.\n"
        "2. 'discord-webhook-url': Credencial de tipo 'Secret text' que resguarda la URL de integración del Webhook de Discord."
    )

    add_image_placeholder(
        "Credenciales Globales registradas en Jenkins",
        "Pega aquí la captura de pantalla de Jenkins (Administrar Jenkins > Credentials > System > Global credentials) "
        "donde se observan 'github-credentials' y 'discord-webhook-url'."
    )

    # -------------------------------------------------------------
    # SECCIÓN 4: REPOSITORIO PRIVADO Y CÓDIGO FUENTE
    # -------------------------------------------------------------
    add_h1("4. Repositorio Privado y Código del Proyecto")
    add_p(
        "Se creó el repositorio privado 'pipeline-ci-jenkins' alojado en la cuenta de GitHub 'AnthonnyM31'. "
        "La estructura del proyecto contiene lógica de negocio modularizada y pruebas unitarias exhaustivas."
    )

    add_image_placeholder(
        "Repositorio Privado en GitHub",
        "Pega aquí la captura de pantalla de la página principal del repositorio en GitHub donde se vea el nombre "
        "'AnthonnyM31 / pipeline-ci-jenkins' y la insignia distintiva de 'Private'."
    )

    add_h2("Código del Pipeline (Jenkinsfile)")
    add_p("A continuación se presenta el código declarativo implementado:")
    add_code_box(
'''pipeline {
    agent any

    environment {
        PROJECT_NAME = 'Pipeline-CI-Individual'
    }

    stages {
        stage('Integración al repo privado') {
            steps {
                echo '=== ETAPA 1: INTEGRACIÓN AL REPOSITORIO PRIVADO ==='
                bat 'git status'
                bat 'git log -1 --pretty=format:"Último commit: %h - %s (%an)"'
            }
        }

        stage('Compilación y test automatizados') {
            steps {
                echo '=== ETAPA 2: COMPILACIÓN Y TESTS AUTOMATIZADOS ==='
                bat 'npm run build'
                bat 'npm test'
            }
        }

        stage('Notificación de ejecución') {
            steps {
                echo '=== ETAPA 3: NOTIFICACIÓN DE ESTADO ==='
                echo 'Preparando y despachando notificación al canal de Discord...'
            }
        }
    }

    post {
        success {
            echo '>>> El pipeline finalizó con ÉXITO. Enviando notificación verde...'
            withCredentials([string(credentialsId: 'discord-webhook-url', variable: 'DISCORD_WEBHOOK')]) {
                bat 'set PIPELINE_STATUS=SUCCESS&& node scripts/notify-discord.js'
            }
        }
        failure {
            echo '>>> El pipeline presentó FALLOS. Enviando notificación roja...'
            withCredentials([string(credentialsId: 'discord-webhook-url', variable: 'DISCORD_WEBHOOK')]) {
                bat 'set PIPELINE_STATUS=FAILURE&& node scripts/notify-discord.js'
            }
        }
    }
}'''
    )

    add_image_placeholder(
        "Configuración del Pipeline en Jenkins",
        "Pega aquí la captura de pantalla de la configuración del Job en Jenkins donde se observe 'Pipeline script from SCM', "
        "la URL del repositorio GitHub, la credencial 'github-credentials' y la rama '*/main'."
    )

    # -------------------------------------------------------------
    # SECCIÓN 5: EVIDENCIAS DE EJECUCIÓN (ÉXITO Y FALLO)
    # -------------------------------------------------------------
    add_h1("5. Evidencias de Ejecución y Validación de Estados")
    add_p(
        "Para dar cumplimiento estricto al requisito de 'Notificación de ejecución exitosa o fallida', "
        "se ejecutaron y documentaron los dos escenarios operacionales:"
    )

    add_h2("5.1. Escenario 1: Ejecución Exitosa (Build #2)")
    add_p(
        "Con el código funcional en el repositorio, Jenkins completó las 3 etapas satisfactoriamente. "
        "Las 13 pruebas unitarias fueron aprobadas con cero fallos ('pass 13, fail 0'). "
        "El bloque post-ejecución ejecutó el script de notificación con estado SUCCESS, enviando la alerta verde a Discord."
    )

    add_image_placeholder(
        "Stage View en Jenkins - Ejecución Exitosa (Verde)",
        "Pega aquí la captura del Stage View de Jenkins donde se observan las 3 etapas en verde (Integración, Compilación y test, Notificación) para el Build #2."
    )

    add_image_placeholder(
        "Consola de Jenkins (Console Output) - Build Exitoso",
        "Pega aquí la captura de la consola de Jenkins mostrando 'npm test: pass 13, fail 0' y 'Finished: SUCCESS'."
    )

    add_h2("5.2. Escenario 2: Ejecución Fallida Provocada (Build #3)")
    add_p(
        "Para verificar la sensibilidad del pipeline ante errores de desarrollo, se indujo un fallo intencional "
        "en el archivo 'tests/math.test.js' alterando el resultado esperado de la función suma: "
        "'assert.strictEqual(sum(2, 3), 999)'. Al disparar el Build #3, Jenkins detuvo la ejecución en la Etapa 2 "
        "al detectar el fallo del test, marcando el build como FAILURE y disparando inmediatamente la notificación roja a Discord."
    )

    add_image_placeholder(
        "Stage View en Jenkins - Ejecución Fallida (Rojo)",
        "Pega aquí la captura del Stage View de Jenkins donde la Etapa 2 se marca en ROJO (FAILED) en el Build #3."
    )

    add_image_placeholder(
        "Consola de Jenkins con el Fallo Detectado",
        "Pega aquí la captura de la consola donde se observe la excepción de la prueba unitaria (AssertionError) y la línea 'Finished: FAILURE'."
    )

    add_h2("5.3. Escenario 3: Restauración a Estado Saludable (Build #4)")
    add_p(
        "Una vez comprobada la detección del fallo, se restauró el código a su estado correcto ('sum(2, 3) === 5'), "
        "se realizó el push correspondiente y se ejecutó el Build #4, regresando el pipeline a su estado 100% verde (SUCCESS)."
    )

    add_h2("5.4. Notificaciones Recibidas en Discord (Evidencia Comparativa)")
    add_p(
        "A continuación se presenta la evidencia del canal de Discord configurado, donde se aprecia la recepción "
        "de ambas alertas automáticas mediante tarjetas visuales (Embeds): la verde correspondiente al éxito "
        "y la roja correspondiente al fallo detectado."
    )

    add_image_placeholder(
        "Canal de Discord con Notificaciones de Éxito y Fallo",
        "Pega aquí la captura de tu canal de Discord donde se vean claramente ambos mensajes: "
        "la tarjeta verde [EXITOSO] y la tarjeta roja [FALLIDO] emitidas por el bot de Jenkins."
    )

    # -------------------------------------------------------------
    # SECCIÓN 6: CONCLUSIONES
    # -------------------------------------------------------------
    add_h1("6. Conclusiones y Aprendizajes")
    add_p(
        "1. Automatización Temprana de Calidad: La integración de pruebas unitarias automatizadas en el pipeline CI "
        "impide que defectos o regresiones introducidas por el equipo lleguen a etapas posteriores del ciclo de vida, "
        "reduciendo los costos de corrección.\n\n"
        "2. Seguridad en Canales de Integración: El uso del gestor de credenciales nativo de Jenkins demostró ser la "
        "técnica adecuada para consumir servicios privados (GitHub PAT) y webhooks sin exponer secretos en el control "
        "de versiones.\n\n"
        "3. Trazabilidad y Notificaciones en Tiempo Real: El feedback inmediato a través de plataformas colaborativas "
        "como Discord optimiza la comunicación del equipo de desarrollo, alertando al instante sobre builds rotos para "
        "su rápida subsanación."
    )

    output_path = r"c:\Users\PC Master\ULTIMO SEMESTRE\PROCESOS DE SOFTWARE\CONFIGURACION PIPEPLINE INDIVIDUAL\Informe_Automatizacion_Pipeline_Jenkins_CI.docx"
    doc.save(output_path)
    print(f"Documento generado exitosamente en: {output_path}")

if __name__ == "__main__":
    create_document()
